"""Synthetic fixtures.

Nothing here copies the operator's sample workbooks or sample CSVs: the builders
below reproduce only the *shape* observed in them (BOM, tab-prefixed cells, the
tab layout, the hidden Graph helper sheets) using made-up values, so the tests
stay offline, small and reviewable.
"""

from __future__ import annotations

import csv
import io
from pathlib import Path

from huawei_ap_report.exports import AP_COLUMNS, SSID_COLUMNS

BOM = "\ufeff"

AP_NUMERIC_COLUMNS = {"AP ID", "STA Quantity", "Total Restart Count", "Poweroff restart count"}
AP_PERCENT_COLUMNS = {"STA Access Failure Ratio", "Logout ratio", "CPU Usage", "Memory Usage"}

#: Columns written into the Detail/Graph workbooks as numbers.
AP_INTEGER_COLUMNS = ("AP ID", "STA Quantity", "Total Restart Count", "Poweroff restart count")
#: Columns written as a bare number with the ``%`` sign removed.
AP_PERCENT_ONLY_COLUMNS = ("STA Access Failure Ratio", "Logout ratio", "CPU Usage", "Memory Usage")

#: Bucket labels for ``__GraphSummary`` column B..H (STA quantity).
STA_BUCKET_LABELS = ("0", "1-5", "6-10", "11-20", "21-30", "31-50", ">50")
#: Bucket labels for the three ratio groups (failure ratio, CPU, memory).
RATIO_BUCKET_LABELS = ("0%", "0-1%", "1-2%", "2-5%", "5-10%", ">10%")
PCT_BUCKET_LABELS = ("0-20%", "20-40%", "40-60%", "60-80%", "80-100%", ">100%")

#: Bucket label columns on the visible Graph sheets (I and the gaps are blank).
VISIBLE_BUCKET_COLUMNS = (
    "B", "C", "D", "E", "F", "G", "H",
    "J", "K", "L", "M", "N", "O",
    "P", "Q", "R", "S", "T", "U",
    "V", "W", "X", "Y", "Z", "AA",
)
#: Matching value columns in ``__GraphSummary`` (no blank-gap columns there).
SUMMARY_BUCKET_COLUMNS = (
    "B", "C", "D", "E", "F", "G", "H",
    "I", "J", "K", "L", "M", "N",
    "O", "P", "Q", "R", "S", "T",
    "U", "V", "W", "X", "Y", "Z",
)

CONNECTED_TITLE = "Wireless SSID Performance & Traffic Report"
CONNECTED_HEADER = list(SSID_COLUMNS)
GRAPH_TITLE = "STA AUTOMATIC PERFORMANCE REPORT"
GRAPH_NOTE = "Paste/replace data in 'Raw Data' — summary tables and charts use Excel formulas."
GRAPH_KPI_LABELS = {"B5": "Total AP Huawei", "F5": "Total STA", "J5": "Avg CPU", "N5": "Avg Memory"}
#: Group 1 is STA quantity, group 2 the failure ratio, groups 3 and 4 the CPU
#: and memory percentage histograms.
GRAPH_SUMMARY_HEADER = (
    ["Report Date"]
    + list(STA_BUCKET_LABELS)
    + list(RATIO_BUCKET_LABELS)
    + list(PCT_BUCKET_LABELS)
    + list(PCT_BUCKET_LABELS)
    + ["Count", "Total STA", "Avg CPU", "Avg Memory"]
)
#: Row 53/54 on the visible Graph tabs, same four groups as the summary sheet.
VISIBLE_BUCKET_LABELS = list(GRAPH_SUMMARY_HEADER[1:-4])


def ap_record(
    ap_id: int = 0,
    name: str = "AP-TEST-01",
    *,
    sta: int = 4,
    failure_ratio: str = "72%",
    logout_ratio: str = "0%",
    cpu: str = "1%",
    memory: str = "58%",
    throughput: str = "537/102",
    login_period: str = "1D:22H:46M:35S",
    restarts: int = 64,
    poweroff_restarts: int = 59,
    status: str = "normal",
) -> dict[str, str]:
    """One AP row in the portal's own vocabulary."""

    values = {
        "AP ID": str(ap_id),
        "AP name": name,
        "Status": status,
        "MAC address": f"6c04-7ac9-{ap_id:04x}",
        "AP group": "TEST-LT1",
        "IP address": f"172.16.24.{100 + (ap_id % 100)}",
        "AP type": "AirEngine6760R-51E",
        "System version": "V200R024C00SPC100",
        "Patch version": "--",
        "Serial Number": f"6R25600136{ap_id:02d}",
        "Power Supply Mode": "802.3 at",
        "Data Link Status": "run",
        "Indoor/Outdoor Channel Set": "Outdoor",
        "Installation location": "",
        "Longitude, Latitude": "",
        "Central AP ID": "",
        "Central AP name": "--",
        "Central AP MAC address": "--",
        "STA Quantity": str(sta),
        "STA Access Failure Ratio": failure_ratio,
        "Logout ratio": logout_ratio,
        "CPU Usage": cpu,
        "Memory Usage": memory,
        "Wired-side throughput(Kbps) ↓↑": throughput,
        "Login period": login_period,
        "Scenario": "",
        "Total Restart Count": str(restarts),
        "Poweroff restart count": str(poweroff_restarts),
    }
    return values


def ssid_record(
    ssid: str = "TEST-SSID",
    *,
    users: int = 55,
    aps: int = 449,
    throughput: str = "7.8M/3.1M",
    frames: str = "321.9M/158.6M",
    retransmission: str = "9.1%(2399/26403)",
    loss: str = "0%",
) -> dict[str, str]:
    return {
        "SSID": ssid,
        "User Quantity": str(users),
        "AP Quantity": str(aps),
        "Valid Throughput (bps) ↓↑": throughput,
        "Frame quantity ↓↑": frames,
        "Downlink retransmission ratio": retransmission,
        "Downlink packet loss ratio": loss,
    }


def _to_csv(columns: tuple[str, ...], records: list[dict[str, str]]) -> str:
    buffer = io.StringIO()
    writer = csv.writer(buffer, lineterminator="\n")
    writer.writerow([f"\t{column}" for column in columns])
    for record in records:
        writer.writerow([f"\t{record.get(column, '')}" for column in columns])
    return BOM + buffer.getvalue()


def ap_csv(records: list[dict[str, str]] | None = None, columns: tuple[str, ...] | None = None) -> bytes:
    return _to_csv(columns or AP_COLUMNS, records or [ap_record()]).encode("utf-8")


def ssid_csv(records: list[dict[str, str]] | None = None, columns: tuple[str, ...] | None = None) -> bytes:
    return _to_csv(columns or SSID_COLUMNS, records or [ssid_record()]).encode("utf-8")


def header_only_csv(columns: tuple[str, ...]) -> bytes:
    """A well-formed export that carries a header and no data rows."""

    buffer = io.StringIO()
    csv.writer(buffer, lineterminator="\n").writerow([f"\t{column}" for column in columns])
    return (BOM + buffer.getvalue()).encode("utf-8")


def write_raw_pair(directory: Path, *, ap: bytes | None = None, ssid: bytes | None = None) -> Path:
    """Create ``output/raw/YYYYMMDD/{apInfo,ssidInfo}.csv`` and return the directory."""

    directory.mkdir(parents=True, exist_ok=True)
    (directory / "apInfo.csv").write_bytes(ap if ap is not None else ap_csv())
    (directory / "ssidInfo.csv").write_bytes(ssid if ssid is not None else ssid_csv())
    return directory


# --------------------------------------------------------------------------
# Template builders
# --------------------------------------------------------------------------


def _connected_sheet(worksheet) -> None:
    worksheet["A2"] = CONNECTED_TITLE
    for index, header in enumerate(CONNECTED_HEADER, start=1):
        worksheet.cell(row=4, column=index, value=f"\t{header}")


def build_connected_template(path: Path, day_tabs: tuple[str, ...] = ("01", "02", "03")) -> Path:
    from openpyxl import Workbook

    workbook = Workbook()
    _connected_sheet(workbook.active)
    workbook.active.title = "Master"
    for tab in day_tabs:
        worksheet = workbook.create_sheet(tab)
        _connected_sheet(worksheet)
    worksheet = workbook.create_sheet("Monthly")
    _connected_sheet(worksheet)
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    return path


def _detail_sheet(worksheet) -> None:
    for index, header in enumerate(AP_COLUMNS, start=1):
        worksheet.cell(row=1, column=index, value=header)


def build_detail_template(path: Path, day_tabs: tuple[str, ...] = ("01", "02", "03")) -> Path:
    from openpyxl import Workbook

    workbook = Workbook()
    _detail_sheet(workbook.active)
    workbook.active.title = "Master"
    for tab in day_tabs:
        worksheet = workbook.create_sheet(tab)
        _detail_sheet(worksheet)
    worksheet = workbook.create_sheet("Monthly")
    _detail_sheet(worksheet)
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    return path


def _graph_sheet(worksheet, *, summary_row: int, summary_columns: tuple[str, ...]) -> None:
    from openpyxl.chart import BarChart, Reference
    from openpyxl.utils import column_index_from_string

    worksheet["A2"] = GRAPH_TITLE
    worksheet["A3"] = GRAPH_NOTE
    for coordinate, label in GRAPH_KPI_LABELS.items():
        worksheet[coordinate] = label
    worksheet["B6"] = f"=__GraphSummary!$AA${summary_row}"
    worksheet["F6"] = f"=__GraphSummary!$AB${summary_row}"
    worksheet["J6"] = f"=__GraphSummary!$AC${summary_row}"
    worksheet["N6"] = f"=__GraphSummary!$AD${summary_row}"

    for column, label in zip(VISIBLE_BUCKET_COLUMNS, VISIBLE_BUCKET_LABELS):
        worksheet[f"{column}53"] = label
    for column, summary_column in zip(VISIBLE_BUCKET_COLUMNS, summary_columns):
        worksheet[f"{column}54"] = f"=__GraphSummary!${summary_column}${summary_row}"

    # The example workbook has one bar chart per histogram; the chart data
    # reference is where the group boundaries are recorded.
    spans = [("B", "H"), ("J", "O"), ("P", "U"), ("V", "AA")]
    for offset, (start, end) in enumerate(spans):
        chart = BarChart()
        chart.title = f"histogram {offset}"
        data = Reference(
            worksheet,
            min_col=column_index_from_string(start),
            max_col=column_index_from_string(end),
            min_row=54,
            max_row=54,
        )
        chart.add_data(data, from_rows=True, titles_from_data=False)
        worksheet.add_chart(chart, f"A{56 + offset * 15}")


def build_graph_template(
    path: Path,
    day_tabs: tuple[str, ...] = ("01", "02", "03"),
    *,
    monthly_summary_row: int = 32,
) -> Path:
    from openpyxl import Workbook

    workbook = Workbook()
    _graph_sheet(workbook.active, summary_row=2, summary_columns=SUMMARY_BUCKET_COLUMNS)
    workbook.active.title = "Master"
    for index, tab in enumerate(day_tabs, start=2):
        worksheet = workbook.create_sheet(tab)
        _graph_sheet(worksheet, summary_row=index + 1, summary_columns=SUMMARY_BUCKET_COLUMNS)
    worksheet = workbook.create_sheet("Monthly")
    _graph_sheet(worksheet, summary_row=monthly_summary_row, summary_columns=SUMMARY_BUCKET_COLUMNS)

    data_sheet = workbook.create_sheet("__GraphData")
    data_sheet.sheet_state = "hidden"
    data_sheet.append(["Report Date", *AP_COLUMNS])
    summary_sheet = workbook.create_sheet("__GraphSummary")
    summary_sheet.sheet_state = "hidden"
    summary_sheet.append(GRAPH_SUMMARY_HEADER)

    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    return path


def build_templates(directory: Path, *, day_tabs: tuple[str, ...] = ("01", "02", "03")) -> dict[str, Path]:
    """Create all three templates and return them keyed by report name."""

    directory.mkdir(parents=True, exist_ok=True)
    return {
        "connected": build_connected_template(directory / "Report_Connected_AP_Huawei.xlsx", day_tabs),
        "detail": build_detail_template(directory / "Report_Detail_AP_Huawei.xlsx", day_tabs),
        "graph": build_graph_template(directory / "Report_Graph_AP_Huawei.xlsx", day_tabs),
    }