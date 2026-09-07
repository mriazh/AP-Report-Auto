"""Report generation tests.

The reducers here were derived from the operator's example workbooks and checked
against the example raw data before being encoded, so these tests pin the
observed behaviour (means, `<1%` == 1, unit-scaled averages, duration averaging)
rather than an invented one.
"""

from __future__ import annotations

from datetime import date
from pathlib import Path

import pytest
from openpyxl import load_workbook

from fixtures import (
    SUMMARY_BUCKET_COLUMNS,
    VISIBLE_BUCKET_COLUMNS,
    ap_csv,
    ap_record,
    build_connected_template,
    build_detail_template,
    build_graph_template,
    ssid_csv,
    ssid_record,
)
from huawei_ap_report.archive import publish_pair
from huawei_ap_report.errors import ReportError, TemplateError
from huawei_ap_report.reports.common import (
    ap_login_period_value,
    ap_mean_value,
    ap_row_values,
    ap_throughput_value,
    mean,
    ssid_monthly_values,
)
from huawei_ap_report.reports.publish import MonthPlan, generate_month, rebuild_from_archive, report_filename
from huawei_ap_report.reports.graph import (
    GRAPH_DATA_SHEET,
    GRAPH_SUMMARY_SHEET,
    _parse_bucket_label,
    build_graph,
)

MONTH = (2026, 9)


def _snapshot(output_root: Path, day: date, *, ap=None, ssid=None):
    return publish_pair(
        output_root,
        day,
        ap_bytes=ap if ap is not None else ap_csv(),
        ssid_bytes=ssid if ssid is not None else ssid_csv(),
    )


# --------------------------------------------------------------------------
# Reducers
# --------------------------------------------------------------------------


def test_mean_ignores_missing_values():
    assert mean([1.0, None, 3.0]) == 2.0
    assert mean([None, None]) is None
    assert mean([]) is None


def test_ap_row_values_converts_metrics_to_numbers():
    record = ap_record(ap_id=3, sta=12, failure_ratio="72%", cpu="13%")
    values = dict(zip(__import__("huawei_ap_report.exports", fromlist=["AP_COLUMNS"]).AP_COLUMNS, ap_row_values(record)))

    assert values["AP ID"] == 3
    assert values["STA Quantity"] == 12
    assert values["CPU Usage"] == 13
    assert values["Installation location"] is None
    assert values["Central AP name"] == "--"


def test_ap_mean_value_averages_numbers_and_keeps_full_precision():
    assert ap_mean_value([1, 2, 4]) == pytest.approx(7 / 3)
    assert ap_mean_value([0, 0, 0]) == 0
    assert ap_mean_value([None, 5]) == 5
    assert ap_mean_value([None, None]) is None


def test_ap_text_columns_take_the_most_recent_value():
    assert ap_mean_value(["normal", "abnormal"]) == "abnormal"
    assert ap_mean_value(["normal", None]) == "normal"
    assert ap_mean_value([None, None]) is None


def test_login_period_is_the_mean_duration_rounded_to_seconds():
    values = ["1D:22H:46M:35S", "5H:43M:38S", "1D:5H:34M:41S"]
    # 168395 + 20618 + 106481 = 295494 s over 3 days = 98498 s = 1D:3H:21M:38S.
    assert ap_login_period_value(values) == "1D:3H:21M:38S"
    assert ap_login_period_value(["--"]) is None


def test_throughput_pair_averages_each_direction():
    # AP wired-side throughput is Kbps: the Monthly tab keeps plain integers.
    assert ap_throughput_value(["537/102", "1013/214"]) == "775/158"
    assert ap_throughput_value(["--", None]) is None


def test_ssid_monthly_uses_sample_semantics():
    # GMF-IoT-style: <1% is read as 1, AP Quantity is the mean rounded.
    records = [
        ssid_record("A", users=55, aps=449, throughput="7.8M/3.1M", frames="321.9M/158.6M", retransmission="9.1%(2399/26403)", loss="<1%(1/13979)"),
        ssid_record("A", users=59, aps=453, throughput="1.4M/499.2K", frames="642.1M/218.4M", retransmission="6.0%(553/9257)", loss="0%"),
    ]
    users, aps, throughput, frames, retransmission, loss = ssid_monthly_values(records)[1:]

    assert users == pytest.approx(57.0)
    assert aps == 451
    # (7.8 + 1.4) / 2 = 4.60M, (3.1 + 0.4992) / 2 = 1.7996M -> 1.80M
    assert throughput == "4.60M/1.80M"
    # (321.9 + 642.1) / 2 = 482.00M, (158.6 + 218.4) / 2 = 188.50M
    assert frames == "482.00M/188.50M"
    assert retransmission == "7.55%"
    # one day is 0% and one is <1% -> (0 + 1) / 2
    assert loss == "0.50%"


def test_ssid_monthly_rejects_empty_input():
    with pytest.raises(Exception):
        ssid_monthly_values([])


# --------------------------------------------------------------------------
# Connected
# --------------------------------------------------------------------------


def test_connected_fills_day_tabs_and_monthly(tmp_path):
    output_root = tmp_path / "output"
    snapshots = {}
    for day, users in ((date(2026, 9, 1), 10), (date(2026, 9, 2), 30)):
        snapshot = _snapshot(output_root, day, ssid=ssid_csv([ssid_record("A", users=users), ssid_record("B")]))
        snapshots[snapshot.day] = snapshot

    template = build_connected_template(tmp_path / "templates" / "Report_Connected_AP_Huawei.xlsx")
    from huawei_ap_report.reports.connected import build_connected

    build_connected(template, snapshots, month=MONTH)
    workbook = load_workbook(template, data_only=False)

    day_one = [workbook["01"].cell(row=row, column=1).value for row in (5, 6, 7)]
    assert day_one == ["A", "B", None]
    assert workbook["01"]["B5"].value == 10
    assert workbook["02"]["B5"].value == 30
    # Day 3 has no snapshot: header stays, data area is empty.
    assert workbook["03"]["A4"].value == "\tSSID"
    assert workbook["03"]["A5"].value is None

    monthly = [[workbook["Monthly"].cell(row=row, column=column).value for column in range(1, 8)] for row in (5, 6)]
    assert monthly[0][0] == "A"
    assert monthly[0][1] == pytest.approx(20.0)
    assert monthly[0][2] == 449
    assert monthly[1][0] == "B"


def test_connected_master_keeps_template_shape(tmp_path):
    output_root = tmp_path / "output"
    snapshot = _snapshot(output_root, date(2026, 9, 1))
    template = build_connected_template(tmp_path / "t.xlsx")
    from huawei_ap_report.reports.connected import build_connected

    build_connected(template, {snapshot.day: snapshot}, month=MONTH)
    workbook = load_workbook(template)

    assert workbook["Master"]["A2"].value == "Wireless SSID Performance & Traffic Report"
    assert workbook["Master"]["A4"].value == "\tSSID"
    assert all(workbook["Master"].cell(row=row, column=1).value is None for row in range(5, 9))


def test_connected_creates_missing_day_tab(tmp_path):
    output_root = tmp_path / "output"
    snapshot = _snapshot(output_root, date(2026, 9, 1))
    template = build_connected_template(tmp_path / "t.xlsx")
    from huawei_ap_report.reports.connected import build_connected

    build_connected(template, {snapshot.day: snapshot}, month=MONTH)
    workbook = load_workbook(template)

    assert {"04", "05", "06", "07", "08", "09"} <= set(workbook.sheetnames)


# --------------------------------------------------------------------------
# Detail
# --------------------------------------------------------------------------


def test_detail_writes_numbers_and_monthly_means(tmp_path):
    output_root = tmp_path / "output"
    snapshots = {}
    for index, day in enumerate((date(2026, 9, 1), date(2026, 9, 2))):
        record = ap_record(ap_id=0, name="AP-1", sta=4 + index * 6, cpu="1%", memory="58%", restarts=64 + index)
        snapshot = _snapshot(output_root, day, ap=ap_csv([record]))
        snapshots[snapshot.day] = snapshot

    template = build_detail_template(tmp_path / "t.xlsx")
    from huawei_ap_report.reports.detail import build_detail

    build_detail(template, snapshots, month=MONTH)
    workbook = load_workbook(template)

    header = [workbook["01"].cell(row=1, column=column).value for column in range(1, 5)]
    assert header == ["AP ID", "AP name", "Status", "MAC address"]
    assert workbook["01"]["A2"].value == 0
    assert workbook["01"]["S2"].value == 4
    assert workbook["01"]["T2"].value == 72
    assert workbook["01"]["O2"].value is None

    # Monthly: STA Quantity mean of 4 and 10, Total Restart Count mean of 64 and 65.
    assert workbook["Monthly"]["S2"].value == pytest.approx(7.0)
    assert workbook["Monthly"]["AA2"].value == pytest.approx(64.5)
    assert workbook["Monthly"]["B2"].value == "AP-1"


def test_detail_monthly_averages_login_period_and_throughput(tmp_path):
    output_root = tmp_path / "output"
    snapshots = {}
    periods = ["1D:22H:46M:35S", "5H:43M:38S"]
    throughput = ["537/102", "1013/214"]
    for index, day in enumerate((date(2026, 9, 1), date(2026, 9, 2))):
        record = ap_record(login_period=periods[index], throughput=throughput[index])
        snapshot = _snapshot(output_root, day, ap=ap_csv([record]))
        snapshots[snapshot.day] = snapshot

    template = build_detail_template(tmp_path / "t.xlsx")
    from huawei_ap_report.reports.detail import build_detail

    build_detail(template, snapshots, month=MONTH)
    workbook = load_workbook(template)

    # (168395 + 20618) / 2 = 94506.5 s -> 94507 s, half away from zero.
    assert workbook["Monthly"]["Y2"].value == "1D:2H:15M:7S"
    assert workbook["Monthly"]["X2"].value == "775/158"


# --------------------------------------------------------------------------
# Graph
# --------------------------------------------------------------------------


def test_graph_writes_hidden_sheets_and_per_day_formulas(tmp_path):
    output_root = tmp_path / "output"
    snapshots = {}
    for day in (date(2026, 9, 1), date(2026, 9, 2)):
        records = [
            ap_record(ap_id=0, name="AP-1", sta=0, failure_ratio="0%", cpu="5%", memory="10%"),
            ap_record(ap_id=1, name="AP-2", sta=3, failure_ratio="80%", cpu="15%", memory="50%"),
            ap_record(ap_id=2, name="AP-3", sta=100, failure_ratio="<1%(1/1000)", cpu="--", memory="--"),
        ]
        snapshot = _snapshot(output_root, day, ap=ap_csv(records))
        snapshots[snapshot.day] = snapshot

    template = build_graph_template(tmp_path / "t.xlsx")
    build_graph(template, snapshots, month=MONTH)
    workbook = load_workbook(template)

    data_sheet = workbook[GRAPH_DATA_SHEET]
    assert data_sheet.sheet_state == "hidden"
    assert data_sheet["A1"].value == "Report Date"
    assert data_sheet.max_row == 1 + 6
    assert data_sheet["A2"].value.date() == date(2026, 9, 1)
    assert data_sheet["C3"].value == "AP-2"

    summary = workbook[GRAPH_SUMMARY_SHEET]
    assert summary["A2"].value.date() == date(2026, 9, 1)
    # STA buckets: 0 -> 1, 1-5 -> 1, >50 -> 1
    assert [summary[f"{column}2"].value for column in SUMMARY_BUCKET_COLUMNS[:7]] == [1, 1, 0, 0, 0, 0, 1]
    # Failure ratio 0%, 80%, <1% (read as 1) -> 0%, 0-1%, >10%
    assert [summary[f"{column}2"].value for column in SUMMARY_BUCKET_COLUMNS[7:13]] == [1, 1, 0, 0, 0, 1]
    # CPU: 5 and 15 land in 0-20%, the "--" row is counted nowhere.
    assert [summary[f"{column}2"].value for column in SUMMARY_BUCKET_COLUMNS[13:19]] == [2, 0, 0, 0, 0, 0]
    assert summary["AA2"].value == 3
    assert summary["AB2"].value == 103
    assert summary["AC2"].value == pytest.approx(10.0)
    assert summary["A4"].value == "Monthly Average"
    assert summary["AA4"].value == pytest.approx(3.0)

    day_one = workbook["01"]
    assert day_one["B6"].value == "=__GraphSummary!$AA$2"
    assert day_one["B54"].value == (
        "=INDEX(__GraphSummary!B:B, MATCH(DATE(2026,9,1), __GraphSummary!$A:$A, 0))"
    )
    assert day_one["J54"].value.startswith("=INDEX(__GraphSummary!I:I")
    # Column I is the blank gap on the visible tabs; AA maps to summary Z.
    assert day_one[f"{VISIBLE_BUCKET_COLUMNS[-1]}54"].value.startswith("=INDEX(__GraphSummary!Z:Z")
    assert day_one["N6"].value == "=__GraphSummary!$AD$2"
    assert workbook["02"]["B54"].value.endswith("MATCH(DATE(2026,9,2), __GraphSummary!$A:$A, 0))")
    assert workbook["Monthly"]["B54"].value == "=ROUND(__GraphSummary!B4, 0)"
    assert workbook["Monthly"]["B6"].value == "=__GraphSummary!$AA$4"
    # Day 3 has no snapshot, so its chart row is cleared rather than left stale.
    assert workbook["03"]["B54"].value is None
    assert workbook["03"]["B6"].value is None


def test_graph_clears_stale_summary_rows_when_days_disappear(tmp_path):
    output_root = tmp_path / "output"
    first = _snapshot(output_root, date(2026, 9, 1))
    second = _snapshot(output_root, date(2026, 9, 2))
    template = build_graph_template(tmp_path / "t.xlsx")

    build_graph(template, {first.day: first, second.day: second}, month=MONTH)
    build_graph(template, {first.day: first}, month=MONTH)

    workbook = load_workbook(template)
    summary = workbook[GRAPH_SUMMARY_SHEET]
    # One day plus the Monthly Average row: the second day's row is gone.
    assert summary["A2"].value.date() == date(2026, 9, 1)
    assert summary["A3"].value == "Monthly Average"
    assert summary["A4"].value is None
    # __GraphData held one row per AP per day; only the remaining day is left.
    assert workbook[GRAPH_DATA_SHEET].max_row == 1 + 1
    assert workbook["02"]["B54"].value is None


def test_parse_bucket_label_reads_plain_hyphen_ranges():
    assert _parse_bucket_label("1-5") == (1.0, 5.0, False)
    assert _parse_bucket_label("6-10") == (6.0, 10.0, False)
    assert _parse_bucket_label("0-1%") == (0.0, 1.0, False)
    assert _parse_bucket_label("0") == (0.0, 0.0, True)
    assert _parse_bucket_label(">50") == (50.0, float("inf"), False)
    assert _parse_bucket_label(">10%") == (10.0, float("inf"), False)
    assert _parse_bucket_label("not a bucket") is None


def test_parse_bucket_label_normalises_locale_range_separators():
    """The real template mangles the separator under some locales.

    ``__GraphSummary`` row 1 of ``Report_Graph_AP_Huawei.xlsx`` carries ``1?5``,
    ``6?10`` and ``0?1%`` where the fixture has a hyphen: ``?`` is ASCII 63,
    plus the tilde and dash variants Excel emits for the range character. None
    of these may reach the "cannot read bucket label" TemplateError.
    """

    for separator in ("?", "~", "\uff5e", "\u223c", "\u2013", "\u2014"):
        assert _parse_bucket_label(f"1{separator}5") == (1.0, 5.0, False), separator
        assert _parse_bucket_label(f"6{separator}10") == (6.0, 10.0, False), separator
        assert _parse_bucket_label(f"0{separator}1%") == (0.0, 1.0, False), separator

    # The exact labels seen in the failing live run.
    assert _parse_bucket_label("1?5") == (1.0, 5.0, False)
    assert _parse_bucket_label("6?10") == (6.0, 10.0, False)
    assert _parse_bucket_label("0?1%") == (0.0, 1.0, False)
    assert _parse_bucket_label("1~5") == (1.0, 5.0, False)
    assert _parse_bucket_label("1～5") == (1.0, 5.0, False)
    assert _parse_bucket_label("1∼5") == (1.0, 5.0, False)
    assert _parse_bucket_label("1–5") == (1.0, 5.0, False)
    assert _parse_bucket_label("1—5") == (1.0, 5.0, False)
    # Absolute buckets and junk stay honest after normalising.
    assert _parse_bucket_label(">50") == (50.0, float("inf"), False)
    assert _parse_bucket_label(">10%") == (10.0, float("inf"), False)
    assert _parse_bucket_label("0") == (0.0, 0.0, True)
    assert _parse_bucket_label("nope?x") is None


def test_graph_builds_when_summary_row_one_uses_query_separators(tmp_path):
    """End-to-end: the real-template separator shape must not raise TemplateError."""

    output_root = tmp_path / "output"
    records = [
        ap_record(ap_id=0, name="AP-1", sta=0, failure_ratio="0%", cpu="5%", memory="10%"),
        ap_record(ap_id=1, name="AP-2", sta=3, failure_ratio="80%", cpu="15%", memory="50%"),
        ap_record(ap_id=2, name="AP-3", sta=100, failure_ratio="<1%(1/1000)", cpu="--", memory="--"),
    ]
    snapshot = _snapshot(output_root, date(2026, 9, 1), ap=ap_csv(records))
    template = build_graph_template(tmp_path / "t.xlsx")

    workbook = load_workbook(template)
    summary_header = workbook[GRAPH_SUMMARY_SHEET]
    for cell in summary_header[1]:
        if isinstance(cell.value, str) and "-" in cell.value:
            cell.value = cell.value.replace("-", "?")
    workbook.save(template)

    build_graph(template, {snapshot.day: snapshot}, month=MONTH)

    workbook = load_workbook(template)
    summary = workbook[GRAPH_SUMMARY_SHEET]
    # Same tally as the hyphen fixture: 0 -> 1, 1-5 -> 1, >50 -> 1, and the
    # failure-ratio group reads 0% / 0-1% / >10% from the mangled labels.
    assert [summary[f"{column}2"].value for column in SUMMARY_BUCKET_COLUMNS[:7]] == [1, 1, 0, 0, 0, 0, 1]
    assert [summary[f"{column}2"].value for column in SUMMARY_BUCKET_COLUMNS[7:13]] == [1, 1, 0, 0, 0, 1]
    assert workbook["01"]["B6"].value == "=__GraphSummary!$AA$2"


def test_graph_template_without_hidden_sheets_is_rejected(tmp_path):
    from openpyxl import Workbook

    workbook = Workbook()
    workbook.active.title = "Master"
    workbook.create_sheet("01")
    path = tmp_path / "broken.xlsx"
    workbook.save(path)

    with pytest.raises(TemplateError, match="__GraphData"):
        build_graph(path, {}, month=MONTH)


# --------------------------------------------------------------------------
# Publication
# --------------------------------------------------------------------------


def test_report_filenames_follow_the_sample_naming():
    assert report_filename(2026, 9, "connected") == "2026_09-Report_Connected_AP_Huawei.xlsx"
    assert report_filename(2026, 10, "detail") == "2026_10-Report_Detail_AP_Huawei.xlsx"
    assert report_filename(2026, 12, "graph") == "2026_12-Report_Graph_AP_Huawei.xlsx"


def _plan(tmp_path, templates_dir):
    return MonthPlan(year=2026, month=9, template_dir=templates_dir, output_root=tmp_path / "output")


def test_generate_month_publishes_all_three_reports(tmp_path):
    from fixtures import build_templates

    templates = build_templates(tmp_path / "config" / "templates")
    output_root = tmp_path / "output"
    snapshots = {}
    for day in (date(2026, 9, 1), date(2026, 9, 2)):
        snapshot = _snapshot(output_root, day)
        snapshots[snapshot.day] = snapshot

    outputs = generate_month(_plan(tmp_path, templates["graph"].parent), snapshots)

    for kind in ("connected", "detail", "graph"):
        assert outputs[kind].exists()
        assert outputs[kind].parent == tmp_path / "output" / "report"
    assert not list((tmp_path / "output" / "report").glob(".staging-*"))


def test_generate_month_leaves_previous_reports_intact_on_failure(tmp_path):
    from fixtures import build_templates

    templates = build_templates(tmp_path / "config" / "templates")
    output_root = tmp_path / "output"
    snapshot = _snapshot(output_root, date(2026, 9, 1))
    plan = _plan(tmp_path, templates["graph"].parent)
    first = generate_month(plan, {snapshot.day: snapshot})
    originals = {kind: path.read_bytes() for kind, path in first.items()}

    # A template whose Graph helper sheets are gone makes the third report fail.
    broken = load_workbook(templates["graph"])
    del broken["__GraphData"]
    broken.save(templates["graph"])

    with pytest.raises(TemplateError, match="__GraphData"):
        generate_month(plan, {snapshot.day: snapshot})

    for kind, path in first.items():
        assert path.read_bytes() == originals[kind], kind
    assert not list((tmp_path / "output" / "report").glob(".staging-*"))


def test_rebuild_from_archive_uses_every_valid_day_of_the_month(tmp_path):
    from fixtures import build_templates

    templates = build_templates(tmp_path / "config" / "templates")
    output_root = tmp_path / "output"
    for day in (date(2026, 9, 1), date(2026, 9, 2)):
        _snapshot(output_root, day)
    _snapshot(output_root, date(2026, 8, 31))

    outputs = rebuild_from_archive(_plan(tmp_path, templates["graph"].parent))

    connected = load_workbook(outputs["connected"])["Monthly"]
    assert connected["A5"].value == "TEST-SSID"
    assert connected["B5"].value == pytest.approx(55.0)


def test_rebuild_without_any_snapshot_reports_clearly(tmp_path):
    from fixtures import build_templates

    templates = build_templates(tmp_path / "config" / "templates")

    with pytest.raises(ReportError, match="no valid raw snapshots"):
        rebuild_from_archive(_plan(tmp_path, templates["graph"].parent))


def test_missing_template_names_the_file_to_place(tmp_path):
    from fixtures import build_detail_template

    build_detail_template(tmp_path / "templates" / "Report_Detail_AP_Huawei.xlsx")

    with pytest.raises(TemplateError, match="no template matching 'connected'"):
        generate_month(_plan(tmp_path, tmp_path / "templates"), {})


def test_templates_are_not_modified_in_place(tmp_path):
    from fixtures import build_templates

    templates = build_templates(tmp_path / "config" / "templates")
    before = {kind: path.read_bytes() for kind, path in templates.items()}
    output_root = tmp_path / "output"
    snapshot = _snapshot(output_root, date(2026, 9, 1))

    generate_month(_plan(tmp_path, templates["graph"].parent), {snapshot.day: snapshot})

    for kind, path in templates.items():
        assert path.read_bytes() == before[kind], kind